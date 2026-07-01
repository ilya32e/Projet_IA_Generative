# -*- coding: utf-8 -*-
"""
Génère un diaporama séparé "Annexes — preuves techniques complémentaires"
pour la soutenance RNCP40875.

Ce fichier est volontairement distinct de Soutenance_3_Projets.pptx : le guide
de préparation recommande un maximum de 20 slides pour les 30 minutes
chronométrées, donc rien ici ne doit s'y ajouter. Ces slides restent
disponibles séparément pour préparer ou mobiliser pendant les 20 minutes
d'échange avec le jury (sécurité réseau, performance détaillée, comparatifs
techniques, démonstration GenAI intégrale).

Sortie : Annexes_Preuves_Techniques.pptx (16:9, même charte graphique).
Usage  : python scripts/generate_annexes_pptx.py
Pré-requis images : lancer d'abord  python scripts/prep_images.py
                     puis générer les charts (pipeline_metrics, benchmark, evalgrid)
"""

import os
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --------------------------------------------------------------------------
# Charte graphique (identique à generate_pptx.py)
# --------------------------------------------------------------------------
NAVY = RGBColor(0x14, 0x2A, 0x47)
NAVY2 = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
TEAL_D = RGBColor(0x1E, 0x73, 0x6A)
GOLD = RGBColor(0xE9, 0xC4, 0x6A)
CORAL = RGBColor(0xE7, 0x6F, 0x51)
SLATE = RGBColor(0x2B, 0x33, 0x40)
GREY = RGBColor(0x5C, 0x66, 0x73)
FAINT = RGBColor(0x8A, 0x93, 0x9F)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xDC, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
FONT_H = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)
IMG_DIR = (Path(__file__).resolve().parents[1] / "assets" / "screenshots")

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

ACCENT_CUR = [TEAL]
_page = {"n": 0}


# --------------------------------------------------------------------------
# Helpers (identiques à generate_pptx.py)
# --------------------------------------------------------------------------
def add_slide():
    return prs.slides.add_slide(BLANK)


def add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, lw=1.0):
    shp = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tb.text_frame.word_wrap = True
    return tb, tb.text_frame


def R(p, text, size, color=SLATE, bold=False, italic=False, font=FONT):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    return run


def footer(slide, tag, accent):
    _page["n"] += 1
    add_rect(slide, Inches(0.55), Inches(7.07), Inches(12.23), Pt(1.2), LINE)
    tb, tf = textbox(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.3))
    R(tf.paragraphs[0], tag, 9, FAINT)
    tb, tf = textbox(slide, Inches(11.6), Inches(7.12), Inches(1.2), Inches(0.3))
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.RIGHT
    R(pp, f"{_page['n']:02d}", 9, accent, bold=True)


def header(slide, kicker, title, accent=TEAL, foot="Annexes · RNCP40875 · Blocs 1 & 2"):
    add_rect(slide, 0, 0, SW, Inches(1.18), NAVY)
    add_rect(slide, 0, Inches(1.18), SW, Pt(3.5), accent)
    add_rect(slide, Inches(0.55), Inches(0.34), Pt(5), Inches(0.52), accent)
    tb, tf = textbox(slide, Inches(0.75), Inches(0.14), Inches(11.8), Inches(1.0))
    R(tf.paragraphs[0], kicker.upper(), 11.5, GOLD, bold=True)
    p2 = tf.add_paragraph()
    R(p2, title, 25, WHITE, bold=True, font=FONT_H)
    footer(slide, foot, accent)


def bullets(tf, items, size=16):
    first = True
    for level, text, bold in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(5)
        p.space_before = Pt(1)
        if level == 0:
            R(p, "▸ ", size, ACCENT_CUR[0], bold=True)
            R(p, text, size, NAVY2 if bold else SLATE, bold=bold)
        else:
            R(p, "•  ", size - 2, FAINT)
            R(p, text, size - 2, GREY)


def bullets_box(slide, items, x, y, w, h, size=16):
    tb, tf = textbox(slide, x, y, w, h)
    bullets(tf, items, size=size)
    return tb


def pic(slide, name, bx, by, bw, bh, frame=True):
    path = IMG_DIR / name
    if not path.exists():
        ph = add_rect(slide, bx, by, bw, bh, LIGHT, line=LINE)
        tb, tf = textbox(slide, bx, by + bh / 2 - Inches(0.2), bw, Inches(0.4))
        pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        R(pp, f"[{name}]", 11, FAINT, italic=True)
        return
    iw, ih = Image.open(path).size
    scale = min(bw / iw, bh / ih)
    w = int(iw * scale); h = int(ih * scale)
    x = int(bx + (bw - w) / 2); y = int(by + (bh - h) / 2)
    off = Emu(45720)
    if frame:
        add_rect(slide, x + off, y + off, w, h, RGBColor(0xC9, 0xD0, 0xDA))
        pad = Emu(28000)
        add_rect(slide, x - pad, y - pad, w + 2 * pad, h + 2 * pad, CARD, line=LINE, lw=1.0)
    p = slide.shapes.add_picture(str(path), x, y, w, h)
    p.line.color.rgb = RGBColor(0xE4, 0xE8, 0xEE)
    p.line.width = Pt(0.5)
    p.shadow.inherit = False
    return p


def img_caption(slide, text, x, y, w):
    tb, tf = textbox(slide, x, y, w, Inches(0.32))
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    R(pp, text, 10, FAINT, italic=True)


def table(slide, rows, top, left, width, height, col_w=None, fs=13, accent=TEAL):
    nr, nc = len(rows), len(rows[0])
    t = slide.shapes.add_table(nr, nc, int(left), int(top), int(width), int(height)).table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = int(cw)
    for r in range(nr):
        for c in range(nc):
            cell = t.cell(r, c)
            cell.margin_left = Inches(0.09); cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            head = (r == 0)
            lines = rows[r][c].split("\n")
            for li, line in enumerate(lines):
                para = cell.text_frame.paragraphs[0] if li == 0 else cell.text_frame.add_paragraph()
                R(para, line, fs + (1 if head else 0),
                  WHITE if head else SLATE, bold=head or c == 0)
            cell.fill.solid()
            if head:
                cell.fill.fore_color.rgb = NAVY
            else:
                cell.fill.fore_color.rgb = CARD if r % 2 else LIGHT
    return t


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ==========================================================================
# SLIDE DE GARDE — ANNEXES
# ==========================================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, NAVY)
add_rect(s, 0, 0, Inches(0.18), SH, GREY)
tb, tf = textbox(s, Inches(0.9), Inches(2.7), Inches(11.4), Inches(0.9))
R(tf.paragraphs[0], "ANNEXES", 40, WHITE, bold=True, font=FONT_H)
tb, tf = textbox(s, Inches(0.9), Inches(3.55), Inches(11.4), Inches(1.2))
R(tf.paragraphs[0],
  "Preuves techniques complémentaires — document séparé, hors des 20 slides / 30 minutes "
  "chronométrées de Soutenance_3_Projets.pptx. À garder ouvert pendant les 20 minutes "
  "d'échange avec le jury (sécurité réseau, performance détaillée, comparatifs techniques, "
  "démonstration GenAI intégrale).", 15, RGBColor(0xC7, 0xD2, 0xDE), italic=True)
notes(s, "Ce document est volontairement séparé de la présentation principale, qui respecte "
         "strictement les vingt slides recommandées par le guide. Ces slides sont prêtes si le "
         "jury pose une question précise sur la sécurité réseau, la performance, le comparatif "
         "de modèles ou la démonstration GenAI.")

# --- Annexe P1.A : sécurité réseau, scalabilité & résilience (détail C1.4) ---
s = add_slide(); header(s, "Annexe · Projet 1", "Sécurité réseau, scalabilité & résilience (détail)", TEAL)
left = [
    (0, "Sécurité des accès & chiffrement", True),
    (1, "authentification JWT OAuth2 (POST /auth/token, Bearer)", False),
    (1, "autorisation : dependency FastAPI sur chaque route sensible", False),
    (1, "secrets : mots de passe stockés en empreinte SHA-256, jamais en clair", False),
    (0, "Limite assumée", True),
    (1, "pas de TLS/HTTPS en local-first (HTTP dev) → à terminer au load-balancer en prod", False),
    (0, "→ C2.1 car accès, authentification et secrets sont traités, pas ignorés", True),
]
right = [
    (0, "Scalabilité — comment l'archi tient la charge", True),
    (1, "nginx devant N réplicas API (docker-compose.lb.yml)", False),
    (1, "testé : --scale api=3 → /metrics.instance alterne entre réplicas (LB actif)", False),
    (0, "Résilience", True),
    (1, "pool_recycle MySQL : reconnexion auto après coupure réseau", False),
    (1, "proxy_next_upstream nginx : bascule si un réplica répond 502/503/504", False),
    (0, "→ C1.4 car la charge et la panne d'un réplica sont gérées, pas subies", True),
]
bullets_box(s, left, Inches(0.7), Inches(1.55), Inches(5.95), Inches(5.3), size=13.5)
bullets_box(s, right, Inches(6.85), Inches(1.55), Inches(5.8), Inches(5.3), size=13.5)
notes(s, "Côté sécurité : authentification JWT OAuth2, autorisation par dependency FastAPI, secrets "
         "jamais en clair (SHA-256). Limite assumée : pas de TLS en local-first. Côté scalabilité : "
         "testé concrètement avec trois réplicas derrière nginx, observable via /metrics.instance. "
         "Résilience à deux niveaux : reconnexion MySQL automatique, failover nginx.")

# --- Annexe P1.B : performance mesurée (détail C2.4) ---
s = add_slide(); header(s, "Annexe · Projet 1", "Performance mesurée — détail (C2.4)", TEAL)
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(11.9), Inches(0.4))
R(tf.paragraphs[0], "Deux instruments de mesure : profiler interne du pipeline ETL  +  benchmark de "
                     "charge de l'API  →  C2.4 car mesuré, pas estimé", 13, TEAL_D, italic=True, bold=True)
pic(s, "p3_pipeline_metrics_t.png", Inches(0.6), Inches(1.95), Inches(5.85), Inches(3.05))
img_caption(s, "PipelineProfiler — temps par étape, run réel d'aujourd'hui (total 146,5 s)",
            Inches(0.6), Inches(5.08), Inches(5.85))
pic(s, "p3_benchmark_chart_t.png", Inches(6.85), Inches(1.95), Inches(5.85), Inches(3.05))
img_caption(s, "scripts/benchmark.py — latence p50/p95 par endpoint, run réel d'aujourd'hui",
            Inches(6.85), Inches(5.08), Inches(5.85))
bullets_box(s, [(0, "Lecture", True),
                (1, "goulot d'étranglement identifié : persist_storage (86,6 s sur 146,5 s)", False)],
            Inches(0.7), Inches(5.55), Inches(5.85), Inches(1.0), size=13)
bullets_box(s, [(0, "Lecture", True),
                (1, "8/9 endpoints sous SLA 300 ms p95 ; carte « bâtiment » hors SLA (≈340 ms)", False)],
            Inches(6.85), Inches(5.55), Inches(5.85), Inches(1.0), size=13)
notes(s, "Le profiler interne chronomètre chaque étape du build : persist_storage domine le temps "
         "total. Le benchmark mesure p50/p95/p99 par endpoint contre un objectif de SLA : huit "
         "endpoints sur neuf respectent les 300 ms, le neuvième est identifié et diagnostiqué.")

# --- Annexe P3.A : choix du modèle — comparatif détaillé ---
s = add_slide(); header(s, "Annexe · Projet 3", "Choix du modèle & de l'approche — détail", GOLD)
table(s, [
    ["Choix retenu", "Alternative écartée", "Pourquoi", "Compétence"],
    ["RAG (retrieval SBERT + contexte structuré + génération)", "Fine-tuning d'un LLM",
     "pas de dataset d'entraînement dédié ; RAG reste à jour sans ré-entraînement", "C5.2"],
    ["Contexte structuré injecté (scores, métiers, lacunes)", "Prompt engineering seul (sans retrieval)",
     "sans données réelles injectées, le LLM généralise ou invente", "C5.2"],
    ["Gemini 2.5-flash (cloud)", "Modèle local unique (flan-t5-small)",
     "meilleure qualité de texte ; le modèle local reste un repli, pas le choix unique", "C5.2"],
    ["SBERT multilingue (paraphrase-multilingual)", "TF-IDF seul",
     "capture le sens, pas que les mots ; TF-IDF gardé comme repli si SBERT indisponible", "C5.1"],
    ["Cache + verrou 1 appel API / profil", "Appel LLM à chaque rafraîchissement",
     "coût et latence maîtrisés, sortie déterministe pour la démo", "C5.2"],
], Inches(1.0), Inches(0.55), Inches(12.25), Inches(5.65),
   col_w=[Inches(3.65), Inches(3.0), Inches(3.8), Inches(1.8)], fs=11.5, accent=GOLD)
notes(s, "Chaque brique a été choisie contre une alternative explicite : RAG plutôt que fine-tuning, "
         "contexte structuré plutôt que prompting seul, Gemini avec repli plutôt que modèle local "
         "unique, SBERT multilingue plutôt que TF-IDF seul.")

# --- Annexe P3.B : démo entrée/sortie réelle (intégrale) ---
s = add_slide(); header(s, "Annexe · Projet 3", "Démo entrée/sortie réelle — capture intégrale", GOLD)
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(11.9), Inches(0.4))
R(tf.paragraphs[0], "Capture réelle d'aujourd'hui : profil P01 (Lina Martin) → retrieval → bio générée "
                     "par Gemini  →  C5.2 car ancré dans des données réelles, pas inventé", 12.5, TEAL_D,
  italic=True, bold=True)
add_rect(s, Inches(0.7), Inches(1.95), Inches(5.85), Inches(4.7), LIGHT, line=LINE)
add_rect(s, Inches(0.7), Inches(1.95), Inches(5.85), Inches(0.45), NAVY2)
tb, tf = textbox(s, Inches(0.85), Inches(1.99), Inches(5.6), Inches(0.4))
R(tf.paragraphs[0], "ENTRÉE — contexte structuré (extrait réel)", 13, WHITE, bold=True)
bullets_box(s, [
    (1, "Profil : Lina Martin → rôle visé : BI Analyst", False),
    (1, "Retrieval SBERT (scores réels) : BI Analyst (0,76) · Data Visualization Specialist (0,71) · Data Engineer Junior (0,70)", False),
    (1, "Forces détectées : visualisation inclusive, nettoyage pandas, dataset exploitable", False),
    (1, "Lacunes détectées : identifier blocs forts/faibles, documenter les transformations", False),
], Inches(0.85), Inches(2.55), Inches(5.55), Inches(3.9), size=12.5)
add_rect(s, Inches(6.78), Inches(1.95), Inches(5.85), Inches(4.7), LIGHT, line=LINE)
add_rect(s, Inches(6.78), Inches(1.95), Inches(5.85), Inches(0.45), TEAL_D)
tb, tf = textbox(s, Inches(6.93), Inches(1.99), Inches(5.6), Inches(0.4))
R(tf.paragraphs[0], "SORTIE — Gemini 2.5-flash (mode=gemini_api, réel)", 13, WHITE, bold=True)
tb, tf = textbox(s, Inches(6.93), Inches(2.55), Inches(5.55), Inches(3.9))
tf.word_wrap = True
R(tf.paragraphs[0],
  "« Lina Martin est une future BI Analyst passionnée par la transformation des données brutes en "
  "insights actionnables. Avec un score global de 0,57, elle excelle dans la préparation de datasets "
  "exploitables et la création de visualisations inclusives. Elle maîtrise le nettoyage de données avec "
  "Pandas et sait expliquer clairement ses analyses, y compris leurs limites. Lina est prête à mettre ses "
  "compétences au service d'une équipe dynamique pour construire des tableaux de bord percutants et "
  "aider à la prise de décision. »", 12.5, SLATE, italic=True)
notes(s, "Démonstration réelle : pipeline exécuté aujourd'hui sur le profil P01. Le mode 'gemini_api' "
         "confirme que Gemini a réellement répondu, pas un repli local ou un template.")

# --- Annexe P3.C : grille d'évaluation GenAI (détail C5.3) ---
s = add_slide(); header(s, "Annexe · Projet 3", "Grille d'évaluation GenAI — détail (C5.3)", GOLD)
pic(s, "p1_evalgrid_t.png", Inches(0.6), Inches(1.55), Inches(6.4), Inches(4.85))
img_caption(s, "scripts/param_sweep.py — reports/param_sweep.csv, run réel d'aujourd'hui",
            Inches(0.6), Inches(6.45), Inches(6.4))
bullets_box(s, [
    (0, "Lecture", True),
    (1, "T=0,0 maximise le score « overall » (0,82), quel que soit top_p → réglage retenu en production", False),
    (1, "anomalie réelle observée : T=0,7/top_p=0,8 a basculé sur le modèle local (sortie de 7 mots) — preuve que la cascade de repli se déclenche vraiment", False),
    (0, "Absence de biais", True),
    (1, "test_ranking_does_not_depend_on_target_role : le classement dépend des compétences, pas du rôle déclaré", False),
    (0, "Ajustement appliqué", True),
    (1, "thinking_budget=0 + budget de tokens ×4 → corrige les réponses tronquées de Gemini 2.5-flash", False),
    (0, "→ C5.3 car réglages ajustés sur preuve mesurée, pas au hasard", True),
], Inches(7.3), Inches(1.55), Inches(5.35), Inches(5.4), size=12.5)
notes(s, "Run réel du script param_sweep : douze combinaisons température/top_p, chacune notée par la "
         "grille d'évaluation. La case rouge est honnête : elle prouve que la cascade de repli se "
         "déclenche vraiment, pas seulement en théorie.")

out = Path(os.environ.get("PPTX_ANNEXES_OUT", Path(__file__).resolve().parents[1] / "Annexes_Preuves_Techniques.pptx"))
prs.save(str(out))
print(f"OK -> {out}  ({len(prs.slides)} slides)")
